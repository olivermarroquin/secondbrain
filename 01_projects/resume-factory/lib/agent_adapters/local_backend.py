from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime, timezone


def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat()


def contains_placeholders(s: str) -> bool:
    return ("[PROPOSED:" in s) or ("[PLACEHOLDER:" in s)


def extract_template_headings(master_docx: Path) -> list[str]:
    try:
        from docx import Document  # type: ignore
    except Exception as e:
        raise RuntimeError(
            "python-docx not available. Run via docgen venv python:\n"
            "  ~/secondbrain/07_system/venvs/docgen/bin/python <script> ...\n"
            f"Import error: {e}"
        )
    doc = Document(str(master_docx))
    headings: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.endswith(":"):
            headings.append(t)
    return headings


def normalize_section(section: str, template_headings: list[str]) -> str:
    s = (section or "").strip().upper()
    for h in template_headings:
        if h.rstrip(":").strip().upper() == s:
            return h
    return section


@dataclass
class LocalContext:
    app_path: Path
    family: str
    company: str
    role_slug: str
    jd_text: str
    job_json: dict
    selection_json: dict
    emphasis_json: dict
    template_dir: Path
    template_headings: list[str]


def build_context(app: Path, load_json, read_text) -> LocalContext:
    pipe = app / "resume_refs" / "resume_pipeline"
    job_json_p = pipe / "job.json"
    sel_json_p = pipe / "selection.json"
    emph_json_p = pipe / "emphasis.json"
    jd_raw_p = app / "jd" / "jd-raw.txt"

    for req in (job_json_p, sel_json_p, emph_json_p, jd_raw_p):
        if not req.exists():
            raise FileNotFoundError(str(req))

    jobj = load_json(job_json_p)
    sel = load_json(sel_json_p)
    emph = load_json(emph_json_p)
    jd = read_text(jd_raw_p)

    try:
        template_dir = Path(sel["selected"]["template_dir"])
    except Exception:
        raise RuntimeError(f"selection.json missing selected.template_dir: {sel_json_p}")

    master_docx = template_dir / "resume-master.docx"
    if not master_docx.exists():
        raise FileNotFoundError(str(master_docx))

    headings: list[str] = []

    family = (jobj.get("app") or {}).get("family") or sel.get("app", {}).get("family") or ""
    company = (jobj.get("app") or {}).get("company") or sel.get("app", {}).get("company") or ""
    role_slug = (jobj.get("app") or {}).get("role_slug") or sel.get("app", {}).get("role_slug") or ""

    if not family:
        family = (sel.get("app") or {}).get("family") or "unknown_family"
    if not company:
        company = (sel.get("app") or {}).get("company") or "unknown_company"
    if not role_slug:
        role_slug = (sel.get("app") or {}).get("role_slug") or "unknown_role"

    return LocalContext(
        app_path=app,
        family=family,
        company=company,
        role_slug=role_slug,
        jd_text=jd,
        job_json=jobj,
        selection_json=sel,
        emphasis_json=emph,
        template_dir=template_dir,
        template_headings=headings,
    )


def propose(ctx: LocalContext, rr_occurrence: int) -> str:
    selected_slug = Path(ctx.template_dir).name

    lines: list[str] = []
    lines.append("# Proposed Resume Changes (AI)")
    lines.append("")
    lines.append(f"- APP: {ctx.app_path}")
    lines.append(f"- FAMILY: {ctx.family}")
    lines.append(f"- SELECTED_TEMPLATE: {selected_slug}")
    lines.append(f"- GENERATED_UTC: {now_utc()}")
    lines.append("")

    n = 1
    lines += [
        f"{n})",
        "SECTION: SUMMARY",
        "CHANGE: REPLACE_SECTION",
        "TO:",
        "- Senior QA Automation Engineer with 7+ years building Java/Selenium automation, CI-driven regression suites, and API/data validation for enterprise systems.",
        "- Strong in risk-based test strategy, SQL validation, and cross-functional Agile delivery (Scrum).",
        "- Known for stabilizing flaky suites, improving coverage on critical workflows, and producing decision-ready QA reporting.",
        "",
    ]
    n += 1

    tech_section = "TECHNICAL SKILL"
    lines += [
        f"{n})",
        f"SECTION: {tech_section}",
        "CHANGE: ADD",
        "TO: API Testing: REST, Postman, response/schema validation, workflow chaining, negative testing.",
        "",
    ]
    n += 1

    lines += [
        f"{n})",
        f"SECTION: {tech_section}",
        "CHANGE: ADD",
        "TO: CI/CD & Execution: GitHub Actions/Jenkins, Dockerized test runs, reports/artifacts, gated regressions.",
        "",
    ]
    n += 1

    lines += [
        f"{n})",
        "SECTION: PROJECT EXPERIENCE",
        "CHANGE: ADD",
        "SUBSECTION: ROLES_AND_RESPONSIBILITIES",
        f"SUBSECTION_OCCURRENCE: {rr_occurrence}",
        "TO: • Added API assertions and SQL validation to reduce regression risk on critical workflows; integrated suites into CI for faster feedback and higher release confidence.",
        "",
    ]

    out = "\n".join(lines).rstrip() + "\n"
    if contains_placeholders(out):
        raise RuntimeError("local_backend propose output contains placeholders — forbidden.")
    return out


def make_patch_item(p: dict, template_headings: list[str]) -> dict:
    ch = (p.get("change") or "").strip().upper()
    if ch not in ("ADD", "REPLACE_SECTION", "DELETE"):
        raise RuntimeError(f"Unsupported CHANGE in proposal {p.get('num')}: {ch}")

    section = p.get("section")
    if not isinstance(section, str) or not section.strip():
        raise RuntimeError(f"Proposal {p.get('num')} missing SECTION")

    item: dict = {"num": int(p.get("num")), "op": ch, "section": normalize_section(section, template_headings)}

    if "subsection" in p:
        item["subsection"] = p["subsection"]
    if "subsection_occurrence" in p:
        item["subsection_occurrence"] = int(p["subsection_occurrence"])

    if ch == "ADD":
        to_text = p.get("to_text")
        if not isinstance(to_text, str) or not to_text.strip():
            raise RuntimeError(f"ADD proposal {p.get('num')} missing TO")
        item["to_text"] = to_text.strip()

    elif ch == "REPLACE_SECTION":
        paras = p.get("to_paragraphs")
        to_text = p.get("to_text")
        if isinstance(paras, list) and len(paras) > 0:
            clean = [str(x).strip() for x in paras if str(x).strip()]
            if not clean:
                raise RuntimeError(f"REPLACE_SECTION proposal {p.get('num')} has empty TO paragraphs")
            item["to_paragraphs"] = clean
        elif isinstance(to_text, str) and to_text.strip():
            item["to_text"] = to_text.strip()
        else:
            raise RuntimeError(f"REPLACE_SECTION proposal {p.get('num')} missing TO content")

    elif ch == "DELETE":
        from_text = p.get("from_text")
        if not isinstance(from_text, str) or not from_text.strip():
            raise RuntimeError(f"DELETE proposal {p.get('num')} missing FROM")
        item["from_text"] = from_text.strip()

    blob = json.dumps(item, ensure_ascii=False)
    if contains_placeholders(blob):
        raise RuntimeError(f"Patch contains placeholders — forbidden (num {item.get('num')}).")

    return item
